"""Portable, metadata-only cohort comparisons with explicit candidate coverage."""

from collections import defaultdict
import copy
import math
import json
import tempfile
from pathlib import Path

from ._io import object_hash, read_json, sha256
from ._identity import compare_identity_evidence
from .results import RAW_FIELDS, compact, percent, raw_export_row, table_text

SUCCESS = {"completed", "evaluated", "reused"}
GROUPS = ("asym_pc", "sym_pt")
SCALES = ("pot", "float")
METHODS = ("rtn", "adaround")
PERCENTILES = (99.9, 99.99, 99.999)
TIE_ORDER = {99.99: 0, 99.9: 1, 99.999: 2}
SHEETS = (
    "Summary",
    "Task_Summary",
    "Model_Best",
    "Core_Raw",
    "Core_Measured",
    "Complete10_Raw",
    "All_History",
    "Coverage",
)
RAW_SHEETS = ("Core_Raw", "Core_Measured", "Complete10_Raw", "All_History")
_DETAIL_FIELDS = """record_id model_id model task dataset metric run_id cohort_id condition
state eligible full_validation samples expected_samples accuracy_percent int8_accuracy_percent
fp32_accuracy_percent decrease_pp recovery_percent scope group scale_family
activation_symmetric weight_granularity weight_group_size activation_scale
weight_scale calibration percentile rounding bias_correction activation_recalibration
entropy_implementation selected_rtn selected_adaround selected_cross
candidate_count expected_candidates original_measurement_id baseline_record_id
source_primary_metric model_source_url model_source_revision checkpoint_sha256
top5_percent AP50_percent AP75_percent metrics_json metric_details_json quantization_json
model_identity code_identity preprocessing_identity data_identity environment_identity parameter_identity
source_result result_sha256 source_table source_table_sha256 verification pairing_status reason""".split()


def read_catalog(path):
    path = Path(path).resolve()
    if not path.is_file():
        raise ValueError(
            f"Not a comparison catalog file: {path}. Pass a catalog.json file, not a folder"
        )
    value = read_json(path)
    if value.get("schema_version") != 1:
        raise ValueError("Unsupported comparison catalog")
    relative = Path(value["history"]["path"])
    if relative.is_absolute() or ".." in relative.parts:
        raise ValueError("History path must be catalog-relative")
    history = (path.parent / relative).resolve()
    if not history.is_relative_to(path.parent):
        raise ValueError("History path escapes catalog")
    if sha256(history) != value["history"]["sha256"]:
        raise ValueError("Comparison history hash changed")
    return value, read_json(history)["records"]


def native(record):
    return percent(record.get("metrics", {}).get(record["primary_metric"]))


def complete(record):
    return bool(
        record["state"] in SUCCESS
        and record.get("verified")
        and record.get("full_validation")
        and record.get("split") == "validation"
        and isinstance(record.get("samples"), int)
        and record["samples"] > 0
        and record.get("samples") == record.get("expected_samples")
        and native(record) is not None
    )


def axes(record):
    q = record.get("quantization")
    if not q:
        return None
    a, w = q.get("activation_symmetric"), q.get("weight_granularity")
    group = (
        "asym_pc"
        if a is False and w == "per_channel"
        else "sym_pt"
        if a is True and w == "per_tensor"
        else None
    )
    pair = (q.get("activation_scale"), q.get("weight_scale"))
    scale = (
        "pot"
        if pair == ("pot_ceil", "pot_ceil")
        else "float"
        if pair == ("float", "float")
        else None
    )
    correction = record.get("corrections", {})
    cal = q.get("calibration", {})
    if (
        group is None
        or scale is None
        or q.get("scope") != "all"
        or cal.get("method", "").lower() != "percentile"
        or cal.get("percentile") not in PERCENTILES
        or q.get("activation_dtype") != "int8"
        or q.get("weight_dtype") != "int8"
        or q.get("activation_granularity") != "per_tensor"
        or q.get("weight_symmetric") is not True
        or q.get("bias_dtype") != "int32"
        or correction.get("bias_correction")
        or correction.get("activation_recalibration")
        or q.get("bias_correction")
    ):
        return None
    return group, scale, "adaround" if q.get("adaround") else "rtn", cal["percentile"]


def pairing_identity(record):
    provenance = record.get("provenance", {})
    # Source-model identity requires explicit binding or code/preprocessing evidence.
    source = provenance.get("model_identity")
    binding = record.get("evidence", {}).get("source_model_identity_binding")
    if binding not in {
        "result.source_model_identity",
        "manifest.model_identity",
    } and not (
        provenance.get("code_identity") and provenance.get("preprocessing_identity")
    ):
        source = None
    environment = provenance.get("environment_identity")
    if environment == object_hash({}):
        environment = None
    return dict(
        source_model_identity=source,
        preprocessing_identity=provenance.get("preprocessing_identity"),
        code_identity=provenance.get("code_identity"),
        data_identity=provenance.get("data_identity"),
        environment_identity=environment,
        evaluation_identity=record.get("pair_identity"),
        model_id=record.get("model_id"),
        metric=record.get("primary_metric"),
        dataset=record.get("dataset"),
        split=record.get("split"),
        samples=record.get("samples"),
        expected_samples=record.get("expected_samples"),
    )


def pairing_status(record, lookup):
    baseline = lookup.get(record.get("baseline_record_id"))
    if not baseline:
        return "insufficient_evidence", "insufficient_evidence:FP32_baseline"
    status, reason = compare_identity_evidence(
        pairing_identity(record), pairing_identity(baseline)
    )
    if status != "passed":
        return status, reason
    if not complete(record) or not complete(baseline):
        return (
            "insufficient_evidence",
            "insufficient_evidence:completed_full_validation",
        )
    if (
        not baseline.get("is_fp32", baseline.get("condition") == "fp32")
        or baseline.get("quantization") is not None
    ):
        return "mismatch", "mismatch:baseline_is_not_FP32"
    if native(baseline) <= 0:
        return (
            "insufficient_evidence",
            "insufficient_evidence:positive_FP32_denominator",
        )
    return "passed", None


def paired(record, lookup):
    return (
        lookup.get(record.get("baseline_record_id"))
        if pairing_status(record, lookup)[0] == "passed"
        else None
    )


def raw_row(record, models, lookup):
    model = models[record["model_id"]]
    q = record.get("quantization") or {}
    cal = q.get("calibration", {})
    correction = record.get("corrections", {})
    ax = axes(record)
    pair_state, pair_reason = pairing_status(record, lookup)
    base = (
        lookup.get(record.get("baseline_record_id")) if pair_state == "passed" else None
    )
    value = native(record) if record["state"] in SUCCESS else None
    fp = native(base) if base else None
    source = model.get("source", {})
    evidence = record.get("evidence", {})
    provenance = record.get("provenance", {})
    return dict(
        record_id=record["record_id"],
        model_id=record["model_id"],
        model=model["name"],
        task=model["task"],
        dataset=record.get("dataset"),
        metric=record["primary_metric"],
        run_id=record["run_id"],
        cohort_id=record["cohort_id"],
        condition=record["condition"],
        state=record["state"],
        eligible=base is not None,
        full_validation=record.get("full_validation", False),
        samples=record.get("samples"),
        expected_samples=record.get("expected_samples"),
        accuracy_percent=value,
        int8_accuracy_percent=value if q else None,
        fp32_accuracy_percent=fp,
        decrease_pp=fp - value if fp is not None and value is not None else None,
        recovery_percent=100 * value / fp
        if fp is not None and value is not None
        else None,
        scope=q.get("scope"),
        group=ax[0] if ax else None,
        scale_family=ax[1] if ax else None,
        activation_symmetric=q.get("activation_symmetric"),
        weight_granularity=q.get("weight_granularity"),
        weight_group_size=q.get("weight_group_size"),
        activation_scale=q.get("activation_scale"),
        weight_scale=q.get("weight_scale"),
        calibration=cal.get("method"),
        percentile=cal.get("percentile"),
        rounding=("adaround" if q.get("adaround") else "rtn")
        if q
        else "fp32"
        if record.get("is_fp32")
        else "unknown",
        bias_correction=bool(
            correction.get("bias_correction") or q.get("bias_correction")
        ),
        activation_recalibration=bool(correction.get("activation_recalibration")),
        entropy_implementation=cal.get("entropy_implementation"),
        selected_rtn=False,
        selected_adaround=False,
        selected_cross=False,
        candidate_count=None,
        expected_candidates=None,
        original_measurement_id=record.get(
            "original_measurement_id", record["record_id"]
        ),
        baseline_record_id=record.get("baseline_record_id"),
        source_primary_metric=record.get("source_primary_metric"),
        model_source_url=source.get("repo") or source.get("url"),
        model_source_revision=source.get("commit") or source.get("framework"),
        checkpoint_sha256=source.get("sha256"),
        top5_percent=percent(record.get("metrics", {}).get("top5")),
        AP50_percent=percent(record.get("metrics", {}).get("AP50")),
        AP75_percent=percent(record.get("metrics", {}).get("AP75")),
        metrics_json=compact(record.get("metrics", {})),
        metric_details_json=compact(record.get("metric_details", {})),
        quantization_json=compact(q) if q else None,
        **{
            key: provenance.get(key)
            for key in (
                "model_identity",
                "code_identity",
                "preprocessing_identity",
                "data_identity",
                "environment_identity",
                "parameter_identity",
            )
        },
        source_result=evidence.get("source_result"),
        result_sha256=evidence.get("result_sha256"),
        source_table=evidence.get("source_table"),
        source_table_sha256=evidence.get("source_table_sha256"),
        verification=evidence.get("verification", "unverified"),
        pairing_status=pair_state,
        reason="; ".join(v for v in (record.get("reason"), pair_reason) if v) or None,
    )


def pick(rows):
    return (
        min(
            rows,
            key=lambda row: (
                -row["accuracy_percent"],
                0 if row["rounding"] == "rtn" else 1,
                TIE_ORDER.get(row["percentile"], 99),
                row["record_id"],
            ),
        )
        if rows
        else None
    )


def core_table(catalog, records, models, lookup, history):
    authority = {run: i for i, run in enumerate(catalog["core_sources"])}
    by_slot = defaultdict(list)
    for record in records:
        if record["run_id"] not in authority or not record.get(
            "core_compatible", False
        ):
            continue
        if not record.get("quantization"):
            if not record.get("is_fp32", record["condition"] == "fp32"):
                continue
            key = (record["model_id"], "fp32")
        else:
            ax = axes(record)
            if ax is None:
                continue
            key = (record["model_id"], *ax)
        by_slot[key].append(record)
    rows, coverage = [], []
    for mid, model in models.items():
        slots = [("fp32",)] + [
            (g, s, method, p)
            for g in GROUPS
            for s in SCALES
            for method in METHODS
            for p in PERCENTILES
        ]
        for slot in slots:
            key = (mid, *slot)
            candidates = by_slot[key]
            eligible = [r for r in candidates if paired(r, lookup)]
            chosen = None
            if eligible:
                priority = min(authority[r["run_id"]] for r in eligible)
                preferred = [r for r in eligible if authority[r["run_id"]] == priority]
                identities = {
                    r.get("original_measurement_id", r["record_id"]) for r in preferred
                }
                if len(identities) > 1:
                    raise ValueError("Conflicting canonical measurements: " + str(key))
                chosen = min(preferred, key=lambda r: r["record_id"])
            if chosen:
                row = copy.deepcopy(history[chosen["record_id"]])
            else:
                record = (
                    min(
                        candidates,
                        key=lambda r: (authority[r["run_id"]], r["record_id"]),
                    )
                    if candidates
                    else None
                )
                row = (
                    copy.deepcopy(history[record["record_id"]])
                    if record
                    else dict(
                        model_id=mid,
                        model=model["name"],
                        task=model["task"],
                        dataset=model["dataset"],
                        metric=model["metric"],
                        state="not_measured",
                        eligible=False,
                        expected_samples=model["expected_samples"],
                        scope="all" if slot != ("fp32",) else None,
                        model_source_url=model.get("source", {}).get("repo")
                        or model.get("source", {}).get("url"),
                        model_source_revision=model.get("source", {}).get("commit")
                        or model.get("source", {}).get("framework"),
                    )
                )
                row.update(
                    accuracy_percent=None, recovery_percent=None, decrease_pp=None
                )
                reason = row.get("reason") or (
                    "not_measured" if not record else "incomplete_or_unmatched_FP32"
                )
                if not record and slot != ("fp32",):
                    historical = [
                        r
                        for r in records
                        if r["model_id"] == mid and axes(r) == slot and complete(r)
                    ]
                    if historical:
                        reason = "historical_only_no_verified_core_equivalence"
                row["reason"] = reason
                coverage.append(
                    dict(
                        model_id=mid,
                        model=model["name"],
                        slot="/".join(map(str, slot)),
                        state=row["state"],
                        reason=reason,
                        candidate_records=len(candidates),
                    )
                )
            if slot != ("fp32",):
                group, scale, method, p = slot
                row.update(
                    group=group,
                    scale_family=scale,
                    rounding=method,
                    percentile=p,
                    activation_symmetric=group == "sym_pt",
                    weight_granularity="per_tensor"
                    if group == "sym_pt"
                    else "per_channel",
                    activation_scale="pot_ceil" if scale == "pot" else "float",
                    weight_scale="pot_ceil" if scale == "pot" else "float",
                    calibration="Percentile",
                    condition=f"{group}_{scale}_{method}_p{p}",
                )
            else:
                row.update(condition="fp32", rounding="fp32")
            rows.append({field: row.get(field) for field in _DETAIL_FIELDS})
    return rows, coverage


def best_rows(core, models):
    best = []
    for mid, model in models.items():
        for group in GROUPS:
            for scale in SCALES:
                available = [
                    r
                    for r in core
                    if r["model_id"] == mid
                    and r["group"] == group
                    and r["scale_family"] == scale
                    and r["eligible"]
                ]
                for method in (*METHODS, "cross"):
                    candidates = [
                        r
                        for r in available
                        if method == "cross" or r["rounding"] == method
                    ]
                    chosen = pick(candidates)
                    expected = 6 if method == "cross" else 3
                    count = len(candidates)
                    selected = dict(chosen) if chosen else None
                    if chosen:
                        chosen["selected_" + method] = True
                    matched_rtn = next(
                        (
                            r
                            for r in available
                            if selected
                            and r["rounding"] == "rtn"
                            and r["percentile"] == selected["percentile"]
                        ),
                        None,
                    )
                    best.append(
                        dict(
                            model_id=mid,
                            model=model["name"],
                            task=model["task"],
                            metric=model["metric"],
                            selection="validation_posthoc",
                            group=group,
                            scale_family=scale,
                            method=method,
                            state="confirmed_posthoc_best"
                            if count == expected
                            else "partial_observed_best"
                            if count
                            else "not_measured",
                            candidate_count=count,
                            expected_candidates=expected,
                            selected_method=selected["rounding"] if selected else None,
                            selected_percentile=selected["percentile"]
                            if selected
                            else None,
                            fp32_accuracy_percent=selected["fp32_accuracy_percent"]
                            if selected
                            else None,
                            accuracy_percent=selected["accuracy_percent"]
                            if selected
                            else None,
                            recovery_percent=selected["recovery_percent"]
                            if selected
                            else None,
                            decrease_pp=selected["decrease_pp"] if selected else None,
                            rtn_at_selected_percentile_record_id=matched_rtn[
                                "record_id"
                            ]
                            if matched_rtn
                            else None,
                            rtn_at_selected_percentile_source_result=matched_rtn[
                                "source_result"
                            ]
                            if matched_rtn
                            else None,
                            rtn_at_selected_percentile_percent=matched_rtn[
                                "accuracy_percent"
                            ]
                            if matched_rtn
                            else None,
                            adaround_same_percentile_change_pp=(
                                selected["accuracy_percent"]
                                - matched_rtn["accuracy_percent"]
                                if selected
                                and selected["rounding"] == "adaround"
                                and matched_rtn
                                else None
                            ),
                            record_id=selected["record_id"] if selected else None,
                            run_id=selected["run_id"] if selected else None,
                            source_result=selected["source_result"]
                            if selected
                            else None,
                        )
                    )
                for row in (
                    r
                    for r in core
                    if r["model_id"] == mid
                    and r["group"] == group
                    and r["scale_family"] == scale
                ):
                    row["candidate_count"] = sum(
                        r["rounding"] == row["rounding"] for r in available
                    )
                    row["expected_candidates"] = 3
    return best


def fixed_rows(catalog, records, models, lookup, history):
    result = []
    fixed = catalog["fixed_run"]
    for mid, model in models.items():
        for scale in SCALES:
            candidates = [
                history[r["record_id"]]
                for r in records
                if r["run_id"] == fixed
                and r["model_id"] == mid
                and axes(r)
                and axes(r)[:2] == ("asym_pc", scale)
                and paired(r, lookup)
            ]
            for method in (*METHODS, "cross"):
                choices = [
                    r
                    for r in candidates
                    if method == "cross" or r["rounding"] == method
                ]
                expected = 2 if method == "cross" else 1
                if len(choices) > expected:
                    raise ValueError("Ambiguous fixed-percentile comparison")
                chosen = pick(choices)
                result.append(
                    dict(
                        model_id=mid,
                        model=model["name"],
                        task=model["task"],
                        metric=model["metric"],
                        selection="previous_fixed_percentile",
                        group="asym_pc",
                        scale_family=scale,
                        method=method,
                        state="confirmed_fixed"
                        if len(choices) == expected
                        else "incomplete",
                        candidate_count=len(choices),
                        expected_candidates=expected,
                        selected_method=chosen["rounding"] if chosen else None,
                        selected_percentile=chosen["percentile"] if chosen else None,
                        fp32_accuracy_percent=chosen["fp32_accuracy_percent"]
                        if chosen
                        else None,
                        accuracy_percent=chosen["accuracy_percent"] if chosen else None,
                        recovery_percent=chosen["recovery_percent"] if chosen else None,
                        record_id=chosen["record_id"] if chosen else None,
                    )
                )
    return result


def averages(best, models, target):
    groups = defaultdict(list)
    for row in best:
        groups[
            (row["selection"], row["group"], row["scale_family"], row["method"])
        ].append(row)
    summary, tasks = [], []
    expected_tasks = defaultdict(list)
    for mid, model in models.items():
        expected_tasks[model["task"]].append(mid)
    for key, rows in groups.items():
        for task in (None, *expected_tasks):
            members = rows if task is None else [r for r in rows if r["task"] == task]
            expected = len(models) if task is None else len(expected_tasks[task])
            done = [
                r
                for r in members
                if r["state"] in ("confirmed_posthoc_best", "confirmed_fixed")
            ]
            valid = len(members) == expected and len(done) == expected
            recovery = (
                math.fsum(r["recovery_percent"] for r in members) / expected
                if valid
                else None
            )
            row = dict(
                selection=key[0],
                group=key[1],
                scale_family=key[2],
                method=key[3],
                task=task or "all",
                metric=members[0]["metric"] if task else "model_recovery",
                completed_models=len(done),
                expected_models=expected,
                completed_candidates=sum(r["candidate_count"] for r in members),
                expected_candidates=sum(r["expected_candidates"] for r in members),
                mean_fp32_accuracy_percent=math.fsum(
                    r["fp32_accuracy_percent"] for r in members
                )
                / expected
                if valid and task
                else None,
                mean_int8_accuracy_percent=math.fsum(
                    r["accuracy_percent"] for r in members
                )
                / expected
                if valid and task
                else None,
                mean_recovery_percent=recovery,
                target_percent=target,
                target_met=recovery >= target if recovery is not None else None,
                state="complete" if valid else "incomplete",
            )
            (summary if task is None else tasks).append(row)
    return summary, tasks


def complete_cohorts(records, models, lookup, history, fixed_run):
    groups = defaultdict(list)
    for record in records:
        if not paired(record, lookup):
            continue
        q = copy.deepcopy(record.get("quantization"))
        policy = "common_scheme"
        if record["run_id"] == fixed_run and q:
            q["calibration"]["percentile"] = "per_model_fixed"
            policy = "per_model_fixed"
        signature = object_hash(
            [
                record["cohort_id"],
                q,
                record.get("corrections", {}),
                record.get("split"),
                policy,
            ]
        )
        groups[signature].append(record)
    rows = []
    for signature, group in groups.items():
        by_model = defaultdict(list)
        for record in group:
            by_model[record["model_id"]].append(record)
        if set(by_model) != set(models):
            continue
        chosen = []
        for mid in models:
            candidates = by_model[mid]
            if (
                len(
                    {
                        r.get("original_measurement_id", r["record_id"])
                        for r in candidates
                    }
                )
                != 1
            ):
                break
            chosen.append(min(candidates, key=lambda r: r["record_id"]))
        else:
            for record in chosen:
                row = dict(history[record["record_id"]])
                row["complete_condition_id"] = signature
                row["percentile_policy"] = (
                    "per_model_fixed"
                    if record["run_id"] == fixed_run and record.get("quantization")
                    else "common_scheme"
                )
                rows.append(row)
    return rows


def build_comparison(catalog, records):
    models = {r["id"]: r for r in catalog["models"]}
    if len(models) != len(catalog["models"]):
        raise ValueError("Duplicate model definitions")
    lookup = {r["record_id"]: r for r in records}
    if len(lookup) != len(records):
        raise ValueError("Duplicate history record ID")
    if any(r["model_id"] not in models for r in records):
        raise ValueError("History contains an unselected model")
    history = {r["record_id"]: raw_row(r, models, lookup) for r in records}
    core, coverage = core_table(catalog, records, models, lookup, history)
    best = best_rows(core, models)
    fixed = fixed_rows(catalog, records, models, lookup, history)
    summary, task_summary = averages(
        best + fixed, models, catalog.get("target_percent", 98.0)
    )
    for record in records:
        if not history[record["record_id"]]["eligible"]:
            coverage.append(
                dict(
                    model_id=record["model_id"],
                    model=models[record["model_id"]]["name"],
                    record_id=record["record_id"],
                    run_id=record["run_id"],
                    state=record["state"],
                    reason=history[record["record_id"]]["reason"],
                )
            )
    return dict(
        Summary=summary,
        Task_Summary=task_summary,
        Model_Best=best + fixed,
        Core_Raw=core,
        Core_Measured=[
            row
            for row in core
            if raw_export_row(row)[
                "fp32_accuracy" if row["rounding"] == "fp32" else "int8_accuracy"
            ]
            is not None
        ],
        Complete10_Raw=complete_cohorts(
            records, models, lookup, history, catalog["fixed_run"]
        ),
        All_History=list(history.values()),
        Coverage=coverage,
    )


def fields_for(name, rows):
    if name in RAW_SHEETS:
        return RAW_FIELDS
    return list(dict.fromkeys(key for row in rows for key in row)) or ["state"]


def _write_comparison(catalog_path, directory, *, xlsx=False):
    catalog, records = read_catalog(catalog_path)
    sheets = build_comparison(catalog, records)
    directory = Path(directory)
    if directory.exists():
        raise FileExistsError("Comparison snapshot already exists; choose a new output")
    if xlsx:
        try:
            from openpyxl import Workbook
            from openpyxl.utils import get_column_letter
        except ImportError as error:
            raise ImportError(
                "XLSX requires the optional 'report' extra: pip install 'opennpu-quant[report]'"
            ) from error
    directory.mkdir(parents=True)
    metadata = dict(
        schema_version=1,
        catalog_sha256=sha256(catalog_path),
        history_sha256=catalog["history"]["sha256"],
        captured_at=catalog["captured_at"],
        selection="validation_posthoc_not_independent",
        target_percent=catalog.get("target_percent", 98),
        sheet_counts={name: len(rows) for name, rows in sheets.items()},
    )
    (directory / "comparison.json").write_text(
        json.dumps(
            dict(metadata=metadata, sheets=sheets),
            ensure_ascii=False,
            separators=(",", ":"),
            allow_nan=False,
        )
        + "\n",
        encoding="utf-8",
    )
    export_sheets = {
        name: [raw_export_row(row) for row in rows] if name in RAW_SHEETS else rows
        for name, rows in sheets.items()
    }
    for name, rows in export_sheets.items():
        fields = fields_for(name, rows)
        for suffix, delimiter, encoding in (
            ("csv", ",", "utf-8-sig"),
            ("tsv", "\t", "utf-8"),
        ):
            (directory / f"{name}.{suffix}").write_text(
                table_text(rows, fields, delimiter=delimiter), encoding=encoding
            )
    if xlsx:
        wb = Workbook()
        wb.remove(wb.active)
        for name, rows in export_sheets.items():
            ws = wb.create_sheet(name)
            fields = fields_for(name, rows)
            ws.append(fields)
            for row_index, row in enumerate(rows, 2):
                values = [row.get(field) for field in fields]
                if any(isinstance(v, str) and len(v) > 32767 for v in values):
                    raise ValueError("XLSX cell exceeds Excel limit: " + name)
                ws.append(values)
                for column, value in enumerate(values, 1):
                    if isinstance(value, str):
                        ws.cell(row_index, column).data_type = "s"
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for i, field in enumerate(fields, 1):
                ws.column_dimensions[get_column_letter(i)].width = min(
                    38, max(14, len(field) + 2)
                )
                if field == "percentile" or field.endswith(
                    ("_accuracy", "_percent", "_pp")
                ):
                    for cells in ws.iter_cols(min_col=i, max_col=i, min_row=2):
                        for cell in cells:
                            cell.number_format = (
                                "0.############" if field == "percentile" else "0.0000"
                            )
        wb.active = wb.sheetnames.index("Core_Measured")
        wb.save(directory / "CORE10_COMPARISON.xlsx")
    fmt = lambda v: "—" if v is None else f"{v:.4f}"
    lines = [
        "# 핵심 10개 Percentile·AdaRound 비교",
        "",
        f"Snapshot: {catalog['captured_at']}",
        "복원율 = 100 × INT8 / 동일 규약 FP32. 전체 평균은 모델별 동일 가중이며 지표 자체를 섞어 평균하지 않는다.",
        "PoT는 Ceil/Ceil, float는 PoT 제약 없는 INT8 scale이다. BC·activation 재보정은 핵심 선택에서 제외한다.",
        "",
        "| 선택 범위 | 구성 | scale | 방법 | 완료 모델 | 평균 복원율 % | 98% |",
        "|---|---|---|---|---:|---:|---|",
    ]
    for row in sheets["Summary"]:
        verdict = (
            "미확정"
            if row["target_met"] is None
            else "충족"
            if row["target_met"]
            else "미달"
        )
        lines.append(
            "| "
            + " | ".join(
                [
                    row["selection"],
                    row["group"],
                    row["scale_family"],
                    row["method"],
                    f"{row['completed_models']}/{row['expected_models']}",
                    fmt(row["mean_recovery_percent"]),
                    verdict,
                ]
            )
            + " |"
        )
    lines += [
        "",
        "## Task별 평균",
        "",
        "| 선택 범위 | 구성 / scale / 방법 | Task | 완료 | FP32 평균 % | INT8 평균 % | 복원율 평균 % |",
        "|---|---|---|---:|---:|---:|---:|",
    ]
    for row in sheets["Task_Summary"]:
        lines.append(
            "| "
            + " | ".join(
                [
                    row["selection"],
                    f"{row['group']} / {row['scale_family']} / {row['method']}",
                    row["task"],
                    f"{row['completed_models']}/{row['expected_models']}",
                    fmt(row["mean_fp32_accuracy_percent"]),
                    fmt(row["mean_int8_accuracy_percent"]),
                    fmt(row["mean_recovery_percent"]),
                ]
            )
            + " |"
        )
    lines += [
        "",
        "## 모델별 사후 최고값",
        "",
        "각 셀은 선택 Percentile, 정확도 %, 복원율 %, 완료 후보 수를 표시한다. 일부 후보만 완료된 행은 현재 관측 최고값이다.",
    ]
    for group in GROUPS:
        for scale in SCALES:
            lines += [
                "",
                f"### {group} / {scale}",
                "",
                "| 모델 | RTN | AdaRound | 모델별 교차 선택 |",
                "|---|---|---|---|",
            ]
            for model in catalog["models"]:
                cells = [model["name"]]
                for method in (*METHODS, "cross"):
                    row = next(
                        r
                        for r in sheets["Model_Best"]
                        if r["selection"] == "validation_posthoc"
                        and r["model_id"] == model["id"]
                        and r["group"] == group
                        and r["scale_family"] == scale
                        and r["method"] == method
                    )
                    label = (
                        f"{row['selected_method']} "
                        if method == "cross" and row["selected_method"]
                        else ""
                    )
                    score = (
                        f"{label}P{row['selected_percentile']}: {fmt(row['accuracy_percent'])} / {fmt(row['recovery_percent'])}"
                        if row["selected_percentile"]
                        else "—"
                    )
                    cells.append(
                        f"{score} ({row['candidate_count']}/{row['expected_candidates']})"
                    )
                lines.append("| " + " | ".join(cells) + " |")
    lines += [
        "",
        "## 읽는 방법",
        "",
        "- validation_posthoc는 세 Percentile 중 사후 최고값이다. RTN/AdaRound 교차 선택은 모델별이며 scale·구성을 넘지 않는다.",
        "- previous_fixed_percentile은 catalog의 fixed_run에 지정한 모델별 고정 Percentile 결과다.",
        "- 후보가 미완료이면 관측 최고값만 표시하고 최종 평균·목표 판정은 비워 둔다. Sym/PT와 Asym/PC는 두 설정 축이 함께 다르다.",
        "- Model_Best에는 선택 Percentile, 원측정 근거, 같은 Percentile RTN 대비 AdaRound 변화가 있다.",
        "- Core_Raw는 250개 후보 슬롯, Complete10_Raw는 10모델 완료 조건, All_History는 부분·실패·취소·재사용 이력이다.",
        "- Core_Measured는 Core_Raw에서 FP32 또는 INT8 정확도가 실제로 측정된 행만 담으며, 정확도가 0인 행도 포함한다.",
        "- Raw CSV·TSV·Excel은 15개 열을 제공한다. scheme의 설정 항목은 ` | `로 구분하고 percentile 값은 별도 숫자 열에 기록한다. 정확도·복원율은 0–100 기준 숫자이며 복원율은 100을 넘을 수 있다.",
        "- 누락·실패한 측정은 빈 셀로 표시한다. 대응 FP32가 없으면 복원율도 빈 셀이다. 상태·실행 ID·검증 근거는 comparison.json과 Coverage에서 확인한다.",
        "- 중복 재사용은 원측정 ID로 연결한다. 같은 이름의 다른 구현이나 전처리를 합치지 않는다.",
        "",
        "## 파일",
        "",
        "- [Excel](CORE10_COMPARISON.xlsx)" if xlsx else "- XLSX는 선택적 출력이다.",
        "- [후보 raw TSV](Core_Raw.tsv) · [측정된 raw TSV](Core_Measured.tsv) · [10모델 완료 raw TSV](Complete10_Raw.tsv) · [전체 이력 TSV](All_History.tsv)",
        "- [Task별 평균](Task_Summary.tsv) · [모델별 최고값](Model_Best.tsv) · [누락·제외 근거](Coverage.tsv)",
        "",
    ]
    (directory / "RESULTS_SUMMARY_KO.md").write_text("\n".join(lines), encoding="utf-8")
    return metadata


def write_comparison(catalog_path, directory, *, xlsx=False):
    directory = Path(directory).resolve()
    if directory.exists():
        raise FileExistsError("Comparison snapshot already exists; choose a new output")
    directory.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(
        prefix=".comparison-", dir=directory.parent
    ) as scratch:
        staged = Path(scratch) / "report"
        metadata = _write_comparison(catalog_path, staged, xlsx=xlsx)
        if directory.exists():
            raise FileExistsError("Comparison output appeared during generation")
        staged.rename(directory)
    return metadata
