import copy
import csv
import importlib.util
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

from opennpu_quant._io import atomic_json, sha256
from opennpu_quant.comparison import (
    GROUPS,
    SCALES,
    METHODS,
    PERCENTILES,
    SHEETS,
    axes,
    build_comparison,
    read_catalog,
    write_comparison,
    RAW_SHEETS,
    raw_export_row,
)
from opennpu_quant.results import normalize_scheme


def fixture(n=10):
    models, records = [], []
    tasks = (
        ["image_classification"] * 6
        + ["object_detection"] * 2
        + ["semantic_segmentation", "face_detection"]
    )
    for i in range(n):
        mid = f"{i:02d}"
        metric = (
            "top1" if i < 6 else "AP" if i < 8 else "mIoU" if i == 8 else "AP_medium"
        )
        models.append(
            dict(
                id=mid,
                name="Model " + mid,
                task=tasks[i],
                metric=metric,
                dataset="data" + mid,
                expected_samples=100,
            )
        )
        base = dict(
            record_id=mid + "_fp32",
            model_id=mid,
            run_id="sweep",
            cohort_id="sweep",
            condition="fp32",
            state="evaluated",
            verified=True,
            full_validation=True,
            split="validation",
            primary_metric=metric,
            metrics={metric: {"value": 0.5 + i * 0.04, "unit": "fraction"}},
            samples=100,
            expected_samples=100,
            dataset="data" + mid,
            quantization=None,
            corrections={},
            provenance=dict(
                model_identity="source-" + mid,
                code_identity="code",
                preprocessing_identity="pre",
                data_identity="data-" + mid,
                environment_identity="environment",
            ),
            pair_identity=mid,
            baseline_record_id=mid + "_fp32",
            core_compatible=True,
            is_fp32=True,
        )
        records.append(base)
        for group in GROUPS:
            for scale in SCALES:
                for method in METHODS:
                    for p in PERCENTILES:
                        r = copy.deepcopy(base)
                        r.update(
                            record_id=f"{mid}_{group}_{scale}_{method}_{p}",
                            condition=f"{group}_{scale}_{method}_{p}",
                            is_fp32=False,
                            quantization=normalize_scheme(
                                dict(
                                    scope="all",
                                    activation_symmetric=group == "sym_pt",
                                    weight_granularity="per_tensor"
                                    if group == "sym_pt"
                                    else "per_channel",
                                    method="Percentile",
                                    percentile=p,
                                    activation_scale="pot_ceil"
                                    if scale == "pot"
                                    else "float",
                                    weight_scale="pot_ceil"
                                    if scale == "pot"
                                    else "float",
                                    adaround={"steps": 10000}
                                    if method == "adaround"
                                    else None,
                                )
                            ),
                        )
                        r["metrics"][metric]["value"] -= (
                            0.01 if method == "rtn" else 0.005
                        )
                        records.append(r)
    catalog = dict(
        schema_version=1,
        models=models,
        core_sources=["sweep"],
        fixed_run="fixed",
        target_percent=98.0,
        captured_at="2026-10-05T00:00:00+00:00",
    )
    return catalog, records


def subset(rows, **keys):
    return [r for r in rows if all(r.get(k) == v for k, v in keys.items())]


def save_fixture(path, catalog, records):
    atomic_json(path / "history.json", dict(records=records))
    catalog = dict(
        catalog, history=dict(path="history.json", sha256=sha256(path / "history.json"))
    )
    atomic_json(path / "catalog.json", catalog)
    return path / "catalog.json"


class ComparisonTests(unittest.TestCase):
    def test_complete_grid_ties_cross_and_same_percentile(self):
        c, records = fixture()
        sheets = build_comparison(c, records)
        self.assertEqual(SHEETS, tuple(sheets))
        self.assertEqual(250, len(sheets["Core_Raw"]))
        rows = subset(
            sheets["Model_Best"],
            selection="validation_posthoc",
            group="asym_pc",
            scale_family="pot",
        )
        self.assertTrue(all(r["selected_percentile"] == 99.99 for r in rows))
        self.assertTrue(
            all(
                r["selected_method"] == "adaround"
                for r in rows
                if r["method"] == "cross"
            )
        )
        self.assertAlmostEqual(
            0.5,
            next(r for r in rows if r["method"] == "adaround")[
                "adaround_same_percentile_change_pp"
            ],
        )
        # A tie favors RTN even when AdaRound has the preferred percentile.
        changed = next(
            r for r in records if r["record_id"] == "00_asym_pc_pot_rtn_99.9"
        )
        changed["metrics"]["top1"]["value"] = 0.495
        cross = subset(
            build_comparison(c, records)["Model_Best"],
            model_id="00",
            group="asym_pc",
            scale_family="pot",
            method="cross",
        )[0]
        self.assertEqual(
            ("rtn", 99.9), (cross["selected_method"], cross["selected_percentile"])
        )

    def test_missing_failed_cancelled_and_denominator_are_not_zero(self):
        c, records = fixture()
        target = next(
            r for r in records if r["record_id"] == "00_asym_pc_pot_adaround_99.999"
        )
        for state in ("pending", "failed", "cancelled_by_user"):
            target["state"] = state
            sheets = build_comparison(c, records)
            best = subset(
                sheets["Model_Best"],
                model_id="00",
                group="asym_pc",
                scale_family="pot",
                method="adaround",
            )[0]
            self.assertEqual(
                (2, 3, "partial_observed_best"),
                (best["candidate_count"], best["expected_candidates"], best["state"]),
            )
            summary = subset(
                sheets["Summary"],
                selection="validation_posthoc",
                group="asym_pc",
                scale_family="pot",
                method="adaround",
            )[0]
            self.assertEqual(9, summary["completed_models"])
            self.assertIsNone(summary["mean_recovery_percent"])
            self.assertIsNone(summary["target_met"])
            raw = next(
                r for r in sheets["Core_Raw"] if r["record_id"] == target["record_id"]
            )
            self.assertIsNone(raw["accuracy_percent"])
        target["state"] = "evaluated"
        target["pair_identity"] = "other protocol"
        self.assertFalse(
            next(
                r
                for r in build_comparison(c, records)["All_History"]
                if r["record_id"] == target["record_id"]
            )["eligible"]
        )
        records[0]["metrics"]["top1"]["value"] = 0
        self.assertEqual(
            0,
            len(
                subset(
                    build_comparison(c, records)["Core_Raw"],
                    model_id="00",
                    eligible=True,
                )
            ),
        )

    def test_task_averages_use_individual_ratios_and_model_weighting(self):
        c, records = fixture()
        s = build_comparison(c, records)
        mean = subset(
            s["Summary"],
            selection="validation_posthoc",
            group="asym_pc",
            scale_family="pot",
            method="rtn",
        )[0]
        expected = (
            sum(100 * (0.49 + i * 0.04) / (0.5 + i * 0.04) for i in range(10)) / 10
        )
        self.assertAlmostEqual(expected, mean["mean_recovery_percent"])
        tasks = subset(
            s["Task_Summary"],
            selection="validation_posthoc",
            group="asym_pc",
            scale_family="pot",
            method="rtn",
        )
        self.assertEqual([6, 2, 1, 1], [r["expected_models"] for r in tasks])
        self.assertNotAlmostEqual(
            sum(r["mean_recovery_percent"] for r in tasks) / 4, expected
        )
        cls = tasks[0]
        self.assertNotAlmostEqual(
            cls["mean_recovery_percent"],
            cls["mean_int8_accuracy_percent"] / cls["mean_fp32_accuracy_percent"] * 100,
        )
        self.assertIsNone(mean["mean_int8_accuracy_percent"])

    def test_corrections_are_history_only_and_partial_train_excluded(self):
        c, records = fixture(1)
        original = records[1]
        for correction in ("bias_correction", "activation_recalibration"):
            r = copy.deepcopy(original)
            r.update(record_id=correction, corrections={correction: True})
            records.append(r)
            self.assertIsNone(axes(r))
        r = copy.deepcopy(original)
        r.update(record_id="train", split="train", samples=5, full_validation=False)
        records.append(r)
        s = build_comparison(c, records)
        self.assertEqual(28, len(s["All_History"]))
        self.assertEqual(25, len(s["Core_Raw"]))
        self.assertFalse(
            next(x for x in s["All_History"] if x["record_id"] == "train")["eligible"]
        )
        q = original["quantization"]
        q["calibration"] = dict(
            method="Entropy", entropy_implementation="centered_kl_v2"
        )
        q["weight_granularity"], q["weight_group_size"] = "per_group", 64
        row = next(
            x
            for x in build_comparison(c, records)["All_History"]
            if x["record_id"] == original["record_id"]
        )
        self.assertEqual("centered_kl_v2", row["entropy_implementation"])
        self.assertEqual(64, row["weight_group_size"])

    def test_reuse_duplicate_cohort_conflict_and_source_priority(self):
        c, records = fixture(1)
        original = records[1]
        original["original_measurement_id"] = "original"
        duplicate = copy.deepcopy(original)
        duplicate.update(record_id="copy", original_measurement_id="original")
        records.append(duplicate)
        s = build_comparison(c, records)
        self.assertEqual(25, len(s["Complete10_Raw"]))
        duplicate["original_measurement_id"] = "independent"
        with self.assertRaisesRegex(ValueError, "Conflicting"):
            build_comparison(c, records)
        duplicate.update(
            run_id="older",
            cohort_id="other",
            metrics={"top1": {"value": 0.99, "unit": "fraction"}},
        )
        c["core_sources"].append("older")
        s = build_comparison(c, records)
        row = subset(
            s["Core_Raw"],
            group="asym_pc",
            scale_family="pot",
            rounding="rtn",
            percentile=99.9,
        )[0]
        self.assertEqual(original["record_id"], row["record_id"])
        self.assertNotEqual(99, row["accuracy_percent"])

    def test_fixed_cohort_does_not_fill_missing_percentiles(self):
        c, records = fixture(1)
        fixed = [
            copy.deepcopy(r)
            for r in records
            if r["is_fp32"]
            or (axes(r) and axes(r)[0] == "asym_pc" and axes(r)[3] == 99.99)
        ]
        for r in fixed:
            r["run_id"] = r["cohort_id"] = "fixed"
            r["record_id"] += "fixed"
            r["baseline_record_id"] += "fixed"
        s = build_comparison(c, records + fixed)
        self.assertEqual(
            6,
            len(
                subset(
                    s["Summary"],
                    selection="previous_fixed_percentile",
                    state="complete",
                )
            ),
        )
        self.assertEqual(
            4, len(subset(s["Complete10_Raw"], percentile_policy="per_model_fixed"))
        )

    def test_csv_json_and_metadata_only_cli_and_hash(self):
        c, records = fixture(1)
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            cp = save_fixture(root, c, records)
            write_comparison(cp, root / "out")
            data = json.loads((root / "out/comparison.json").read_text())
            for name in SHEETS:
                for suffix, delimiter in (("csv", ","), ("tsv", "\t")):
                    with (root / f"out/{name}.{suffix}").open(
                        encoding="utf-8-sig"
                    ) as stream:
                        rows = list(csv.DictReader(stream, delimiter=delimiter))
                    self.assertEqual(len(data["sheets"][name]), len(rows))
                    for x, y in zip(rows, data["sheets"][name]):
                        if name in RAW_SHEETS:
                            y = raw_export_row(y)
                        for key, value in y.items():
                            if isinstance(value, (int, float)) and not isinstance(
                                value, bool
                            ):
                                self.assertEqual(value, float(x[key]))
            code = "from opennpu_quant.cli import main;import sys;main(['compare',sys.argv[1],'--output',sys.argv[2]]);assert not set(['numpy','onnx','onnxruntime','torch','tensorflow','cv2']) & set(sys.modules)"
            subprocess.run(
                [sys.executable, "-c", code, str(cp), str(root / "cli")],
                check=True,
                stdout=subprocess.DEVNULL,
            )
            with self.assertRaises(FileExistsError):
                write_comparison(cp, root / "out")
            (root / "history.json").write_text("{}")
            with self.assertRaisesRegex(ValueError, "hash"):
                read_catalog(cp)

    @unittest.skipUnless(importlib.util.find_spec("openpyxl"), "optional XLSX extra")
    def test_xlsx_numeric_blank_and_formula_safety(self):
        from openpyxl import load_workbook

        c, records = fixture(1)
        c["models"][0]["name"] = "=1+1"
        records[1]["state"] = "failed"
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            cp = save_fixture(root, c, records)
            write_comparison(cp, root / "out", xlsx=True)
            wb = load_workbook(root / "out/CORE10_COMPARISON.xlsx", read_only=True)
            self.assertEqual(list(SHEETS), wb.sheetnames)
            ws = wb["Core_Raw"]
            rows = list(ws)
            columns = {cell.value: i for i, cell in enumerate(rows[0])}
            for cells in rows[1:]:
                self.assertEqual("s", cells[columns["model"]].data_type)
                acc = cells[columns["int8_accuracy"]]
                if acc.value is not None:
                    self.assertEqual("n", acc.data_type)
            self.assertIsNone(rows[2][columns["int8_accuracy"]].value)
            wb.close()
