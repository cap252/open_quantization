"""Spreadsheet contracts for comparison raw exports."""

import copy
import csv
import importlib.util
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

from opennpu_quant._io import atomic_json, object_hash, sha256
from opennpu_quant.comparison import (
    build_comparison,
    raw_export_row,
    write_comparison,
)
from opennpu_quant.results import export_records, normalize_scheme, tables


HEADERS = """model task dataset metric fp32_accuracy int8_accuracy recovery_percent
scheme percentile model_source_url model_source_revision top5_percent AP50_percent AP75_percent
metrics_json""".split()
RAW_SHEETS = ("Core_Raw", "Core_Measured", "Complete10_Raw", "All_History")


def fixture():
    catalog = dict(
        schema_version=1,
        models=[
            dict(
                id="toy",
                name="테스트 모델",
                task="object_detection",
                dataset="coco",
                metric="AP",
                expected_samples=100,
                source=dict(repo="https://example.test/model", commit="revision-1"),
            )
        ],
        core_sources=["sweep"],
        fixed_run="fixed",
        captured_at="2026-10-06T00:00:00Z",
    )
    base = dict(
        record_id="fp32",
        model_id="toy",
        run_id="sweep",
        cohort_id="sweep",
        condition="fp32",
        state="evaluated",
        verified=True,
        full_validation=True,
        split="validation",
        primary_metric="AP",
        metrics={"AP": {"value": 0.8, "unit": "fraction"}},
        samples=100,
        expected_samples=100,
        dataset="coco",
        quantization=None,
        provenance=dict(
            model_identity="source",
            code_identity="code",
            preprocessing_identity="pre",
            data_identity="data",
            environment_identity="env",
        ),
        pair_identity="same-evaluation",
        baseline_record_id="fp32",
        core_compatible=True,
        is_fp32=True,
    )
    quantized = copy.deepcopy(base)
    quantized.update(
        record_id="int8",
        condition="pot_adaround",
        is_fp32=False,
        quantization=normalize_scheme(
            dict(
                scope="all",
                activation_symmetric=False,
                weight_granularity="per_channel",
                activation_scale="pot_ceil",
                weight_scale="pot_ceil",
                method="percentile",
                percentile=99.99,
                adaround={"steps": 10000},
            )
        ),
        metrics={
            "AP": {"value": 78, "unit": "percent"},
            "top5": {"value": 0.95, "unit": "fraction"},
            "AP50": {"value": 90, "unit": "percent"},
            "AP75": {"value": 0.7, "unit": "fraction"},
        },
    )
    return catalog, [base, quantized]


class RawExportTests(unittest.TestCase):
    def setUp(self):
        folder = tempfile.TemporaryDirectory()
        self.addCleanup(folder.cleanup)
        self.root = Path(folder.name)
        self.catalog, self.records = fixture()

    def save_catalog(self):
        atomic_json(self.root / "history.json", dict(records=self.records))
        catalog = dict(
            self.catalog,
            history=dict(
                path="history.json", sha256=sha256(self.root / "history.json")
            ),
        )
        atomic_json(self.root / "catalog.json", catalog)
        return self.root / "catalog.json"

    def history_rows(self):
        sheets = build_comparison(self.catalog, self.records)
        return [raw_export_row(row) for row in sheets["All_History"]]

    def test_columns_units_and_source_values(self):
        base, quantized = self.history_rows()
        self.assertEqual(HEADERS, list(quantized))
        self.assertEqual("테스트 모델", quantized["model"])
        self.assertEqual("object_detection", quantized["task"])
        self.assertEqual("coco", quantized["dataset"])
        self.assertEqual("AP", quantized["metric"])
        self.assertEqual(80, quantized["fp32_accuracy"])
        self.assertEqual(78, quantized["int8_accuracy"])
        self.assertEqual(97.5, quantized["recovery_percent"])
        self.assertEqual([95, 90, 70], [quantized[k] for k in HEADERS[11:14]])
        self.assertEqual(99.99, quantized["percentile"])
        self.assertEqual("https://example.test/model", quantized["model_source_url"])
        self.assertEqual("revision-1", quantized["model_source_revision"])
        self.assertEqual(
            self.records[1]["metrics"], json.loads(quantized["metrics_json"])
        )
        self.assertEqual("FP32", base["scheme"])
        self.assertEqual(80, base["fp32_accuracy"])
        self.assertIsNone(base["int8_accuracy"])
        self.assertIsNone(base["percentile"])
        self.assertEqual(
            "scope=all | activation=asymmetric | weight=per_channel | "
            "activation_scale=pot_ceil | weight_scale=pot_ceil | "
            "calibration=percentile | rounding=AdaRound",
            quantized["scheme"],
        )

    def test_unmeasured_candidates_keep_distinct_scheme_percentile_pairs(self):
        sheets = build_comparison(self.catalog, [])
        rows = [raw_export_row(row) for row in sheets["Core_Raw"]]
        self.assertEqual(25, len(rows))
        self.assertEqual(9, len({row["scheme"] for row in rows}))
        self.assertEqual(25, len({(row["scheme"], row["percentile"]) for row in rows}))
        self.assertEqual([], sheets["Core_Measured"])
        self.assertEqual("FP32", rows[0]["scheme"])
        for row in rows:
            self.assertEqual(HEADERS, list(row))
            for key in ("fp32_accuracy", "int8_accuracy", "recovery_percent"):
                self.assertIsNone(row[key])
            self.assertEqual("revision-1", row["model_source_revision"])
        self.records = []
        write_comparison(self.save_catalog(), self.root / "empty")
        for name in RAW_SHEETS:
            with (self.root / f"empty/{name}.tsv").open() as stream:
                self.assertEqual(HEADERS, next(csv.reader(stream, delimiter="\t")))

    def test_scale_modes_granularity_and_corrections_remain_distinct(self):
        q = self.records[1]["quantization"]
        original = copy.deepcopy(q)
        labels = {self.history_rows()[1]["scheme"]}
        for key, values in (
            ("scope", ["basic"]),
            ("activation_symmetric", [True]),
            ("weight_granularity", ["per_tensor", "per_group"]),
            ("activation_scale", ["float", "pot_nearest"]),
            ("weight_scale", ["float", "pot_nearest", "pot_mse"]),
            ("adaround", [None]),
        ):
            for value in values:
                q.clear()
                q.update(copy.deepcopy(original))
                q[key] = value
                label = self.history_rows()[1]["scheme"]
                self.assertNotIn(label, labels)
                self.assertTrue(all("=" in part for part in label.split(" | ")))
                labels.add(label)
        q.update(original)
        q["weight_granularity"] = "per_group"
        q["weight_group_size"] = 64
        q["calibration"] = dict(
            method="Entropy", entropy_implementation="centered_kl_v2"
        )
        self.records[1]["corrections"] = dict(
            bias_correction=True, activation_recalibration=True
        )
        label = self.history_rows()[1]["scheme"]
        for part in (
            "calibration=entropy",
            "weight_group_size=64",
            "entropy_implementation=centered_kl_v2",
            "bias_correction=true",
            "activation_recalibration=true",
        ):
            self.assertIn(part, label.split(" | "))
        self.assertNotIn("percentile=", label)

    def test_failed_unmatched_zero_and_unknown_units_are_not_fabricated(self):
        for state in ("pending", "failed", "cancelled_by_user"):
            with self.subTest(state=state):
                self.records[1]["state"] = state
                row = self.history_rows()[1]
                for key in ("int8_accuracy", "recovery_percent", *HEADERS[11:14]):
                    self.assertIsNone(row[key])
                self.assertEqual("{}", row["metrics_json"])
        self.records[1]["state"] = "evaluated"
        self.records[1]["provenance"]["preprocessing_identity"] = "different"
        row = self.history_rows()[1]
        self.assertEqual(78, row["int8_accuracy"])
        self.assertIsNone(row["fp32_accuracy"])
        self.assertIsNone(row["recovery_percent"])
        self.records[1]["provenance"]["preprocessing_identity"] = "pre"
        self.records[0]["metrics"]["AP"]["value"] = 0
        self.assertEqual(0, self.history_rows()[0]["fp32_accuracy"])
        self.assertIsNone(self.history_rows()[1]["recovery_percent"])
        self.records[1]["metrics"]["AP75"]["unit"] = "unknown"
        self.assertIsNone(self.history_rows()[1]["AP75_percent"])

    def test_csv_tsv_schema_formula_safety_and_json_evidence(self):
        self.catalog["models"][0]["name"] = "=1+1\t모델\n줄"
        write_comparison(self.save_catalog(), self.root / "out")
        data = json.loads((self.root / "out/comparison.json").read_text())
        for name in RAW_SHEETS:
            exported = []
            for suffix, delimiter in (("csv", ","), ("tsv", "\t")):
                path = self.root / f"out/{name}.{suffix}"
                if suffix == "csv":
                    self.assertTrue(path.read_bytes().startswith(b"\xef\xbb\xbf"))
                with path.open(encoding="utf-8-sig", newline="") as stream:
                    reader = csv.DictReader(stream, delimiter=delimiter)
                    self.assertEqual(HEADERS, reader.fieldnames)
                    rows = list(reader)
                self.assertEqual(len(data["sheets"][name]), len(rows))
                self.assertTrue(all(row["model"].startswith("'=1") for row in rows))
                exported.append(rows)
            self.assertEqual(*exported)
        evidence = data["sheets"]["All_History"][1]
        self.assertEqual("passed", evidence["pairing_status"])
        self.assertEqual("int8", evidence["record_id"])
        self.assertTrue(evidence["eligible"])

    @unittest.skipUnless(importlib.util.find_spec("openpyxl"), "optional XLSX extra")
    def test_excel_has_only_requested_columns_and_numeric_cells(self):
        from openpyxl import load_workbook

        self.catalog["models"][0]["name"] = "=1+1"
        write_comparison(self.save_catalog(), self.root / "out", xlsx=True)
        workbook = load_workbook(self.root / "out/CORE10_COMPARISON.xlsx")
        self.addCleanup(workbook.close)
        for name in RAW_SHEETS:
            sheet = workbook[name]
            self.assertEqual(HEADERS, [cell.value for cell in sheet[1]])
            self.assertEqual(15, sheet.max_column)
            self.assertEqual("A2", sheet.freeze_panes)
            self.assertEqual(sheet.dimensions, sheet.auto_filter.ref)
            for row in sheet.iter_rows(min_row=2):
                self.assertEqual("s", row[0].data_type)
                for index in (4, 5, 6, 11, 12, 13):
                    self.assertEqual("0.0000", row[index].number_format)
                    if row[index].value is not None:
                        self.assertEqual("n", row[index].data_type)
                self.assertEqual("0.############", row[8].number_format)
                if row[8].value is not None:
                    self.assertEqual("n", row[8].data_type)
        self.assertEqual("Core_Measured", workbook.active.title)
        values = list(workbook["All_History"].values)
        self.assertIsNone(values[1][5])
        self.assertEqual((80, 78, 97.5), values[2][4:7])
        self.assertIsNone(values[1][8])
        self.assertEqual(99.99, values[2][8])

    def test_compare_cli_remains_metadata_only_and_checks_history_hash(self):
        catalog = self.save_catalog()
        result = subprocess.run(
            [
                sys.executable,
                "-c",
                "from opennpu_quant.cli import main; import sys; "
                "main(['compare', sys.argv[1], '--output', sys.argv[2]]); "
                "assert not {'numpy', 'onnx', 'onnxruntime', 'torch', 'tensorflow'} & set(sys.modules)",
                str(catalog),
                str(self.root / "cli"),
            ],
            cwd=self.root,
            text=True,
            capture_output=True,
        )
        self.assertEqual(0, result.returncode, result.stderr)
        with self.assertRaises(FileExistsError):
            write_comparison(catalog, self.root / "cli")
        (self.root / "history.json").write_text("{}")
        with self.assertRaisesRegex(ValueError, "hash"):
            write_comparison(catalog, self.root / "invalid")
        self.assertFalse((self.root / "invalid").exists())

    def run_records(self):
        return [
            dict(
                record,
                model="toy",
                model_name="테스트 모델",
                origin="standalone_run",
                complete=True,
                evaluation_identity="same-evaluation",
                model_metadata=self.catalog["models"][0],
                model_source=self.catalog["models"][0]["source"],
                evidence=dict(model_identity="source", preprocessing_identity="pre"),
            )
            for record in self.records
        ]

    def test_run_exports_and_both_cli_views_use_the_same_raw_schema(self):
        records = self.run_records()
        export_records(records, self.root / "run")
        exported = []
        for name in ("accuracy", "results"):
            for suffix, delimiter in (("csv", ","), ("tsv", "\t")):
                with (self.root / f"run/{name}.{suffix}").open(
                    encoding="utf-8-sig", newline=""
                ) as stream:
                    reader = csv.DictReader(stream, delimiter=delimiter)
                    self.assertEqual(HEADERS, reader.fieldnames)
                    exported.append(list(reader))
        self.assertTrue(all(rows == exported[0] for rows in exported))
        row = exported[0][1]
        self.assertEqual("object_detection", row["task"])
        self.assertEqual("테스트 모델", row["model"])
        self.assertEqual(97.5, float(row["recovery_percent"]))
        self.assertEqual(self.history_rows()[1]["scheme"], row["scheme"])
        self.assertEqual(99.99, float(row["percentile"]))
        for view in ("accuracy", "full"):
            result = subprocess.run(
                [
                    sys.executable,
                    "-m",
                    "opennpu_quant",
                    "results",
                    str(self.root / "run/records.json"),
                    "--stdout",
                    "tsv",
                    "--view",
                    view,
                ],
                cwd=self.root,
                text=True,
                capture_output=True,
            )
            self.assertEqual(0, result.returncode, result.stderr)
            self.assertEqual(
                HEADERS, next(csv.reader(result.stdout.splitlines(), delimiter="\t"))
            )
            self.assertEqual((self.root / "run/results.tsv").read_text(), result.stdout)
        saved = json.loads((self.root / "run/records.json").read_text())
        self.assertEqual(records, saved["records"])

    def test_run_scheme_marks_activation_recalibration_without_extra_columns(self):
        records = self.run_records()
        records[1]["comparison"] = dict(
            family="adaround",
            activation_recalibration=True,
            evaluation_batch_size=1,
            optimization="ORT_DISABLE_ALL",
            selection_basis="fixed",
        )
        row = raw_export_row(tables(records)[0][1])
        self.assertEqual(HEADERS, list(row))
        self.assertIn("activation_recalibration=true", row["scheme"].split(" | "))

    def test_report_rebuild_and_stdout_use_raw_columns(self):
        policy = dict(
            scope="all",
            activation_symmetric=False,
            weight_granularity="per_channel",
            method="percentile",
            percentile=99.99,
            activation_scale="pot_ceil",
            weight_scale="pot_ceil",
        )
        config = dict(
            name="sweep",
            models=["toy"],
            conditions={
                "fp32": {},
                "pot_adaround": {"scheme": "pot", "adaround": True},
            },
            schemes={"pot": policy},
            adaround={"steps": 10000},
            evaluation={"split": "validation"},
        )
        directory = self.root / "report"
        atomic_json(directory / "effective_config.json", config)
        for record in self.records:
            path = directory / "toy" / record["condition"] / "result.json"
            measured = dict(
                config_identity=object_hash(config),
                model_metadata=self.catalog["models"][0],
                title="테스트 모델",
                metric="AP",
                metrics={k: v["value"] for k, v in record["metrics"].items()},
                metric_units={k: v["unit"] for k, v in record["metrics"].items()},
                samples=100,
                expected_samples=100,
                complete=True,
                model_identity="source",
                preprocessing_identity="pre",
                evaluation_identity="same-evaluation",
                scheme=policy if record["quantization"] else None,
            )
            atomic_json(path, measured)
            atomic_json(path.parent / "complete.json", dict(result_sha256=sha256(path)))
        result = subprocess.run(
            [
                sys.executable,
                "-m",
                "opennpu_quant",
                "report",
                str(directory),
                "--format",
                "csv,tsv,json",
                "--stdout",
                "tsv",
                "--view",
                "full",
            ],
            cwd=self.root,
            text=True,
            capture_output=True,
        )
        self.assertEqual(0, result.returncode, result.stderr)
        rows = list(csv.DictReader(result.stdout.splitlines(), delimiter="\t"))
        self.assertEqual(HEADERS, list(rows[0]))
        self.assertEqual("object_detection", rows[1]["task"])
        self.assertEqual(97.5, float(rows[1]["recovery_percent"]))
        self.assertEqual(99.99, float(rows[1]["percentile"]))
        self.assertEqual((directory / "accuracy.tsv").read_text(), result.stdout)
        self.assertEqual((directory / "results.tsv").read_text(), result.stdout)

    def test_percentile_is_separate_and_other_calibration_methods_stay_distinct(self):
        labels = set()
        for percentile in (99.9, 99.99, 99.999):
            self.records[1]["quantization"]["calibration"]["percentile"] = percentile
            comparison_row = self.history_rows()[1]
            run_row = raw_export_row(tables(self.run_records())[0][1])
            for row in (comparison_row, run_row):
                self.assertEqual(percentile, row["percentile"])
                self.assertIn("calibration=percentile", row["scheme"].split(" | "))
                self.assertNotIn("percentile=", row["scheme"])
                labels.add(row["scheme"])
        self.assertEqual(1, len(labels))
        for method in ("MinMax", "Entropy"):
            self.records[1]["quantization"]["calibration"]["method"] = method
            for row in (
                self.history_rows()[1],
                raw_export_row(tables(self.run_records())[0][1]),
            ):
                self.assertIsNone(row["percentile"])
                self.assertIn("calibration=" + method.lower(), row["scheme"])
                self.assertNotIn(row["scheme"], labels)

    def test_core_measured_preserves_zero_and_unmatched_measurements(self):
        for state in ("pending", "failed", "cancelled_by_user"):
            self.records[1]["state"] = state
            measured = build_comparison(self.catalog, self.records)["Core_Measured"]
            self.assertEqual(["fp32"], [row["record_id"] for row in measured])
        self.records[1]["state"] = "evaluated"
        self.records[1]["metrics"]["AP"]["value"] = 0
        for unmatched in (False, True):
            with self.subTest(unmatched=unmatched):
                if unmatched:
                    self.records[0]["metrics"]["AP"]["value"] = 0
                    self.records[1]["provenance"]["preprocessing_identity"] = (
                        "different"
                    )
                sheets = build_comparison(self.catalog, self.records)
                measured = sheets["Core_Measured"]
                self.assertEqual(
                    ["fp32", "int8"], [row["record_id"] for row in measured]
                )
                self.assertEqual(
                    [row for row in sheets["Core_Raw"] if row["record_id"]], measured
                )
                self.assertEqual(0, raw_export_row(measured[1])["int8_accuracy"])
                if unmatched:
                    self.assertEqual(0, raw_export_row(measured[0])["fp32_accuracy"])
                    self.assertIsNone(raw_export_row(measured[1])["recovery_percent"])


if __name__ == "__main__":
    unittest.main()
